"""
AI Due Diligence Pipeline
=========================
8-Phase, 30-Prompt chain for generating IC-grade analysis from a CIM.
Each phase output feeds into subsequent phases via a variable store.

Usage:
    from ai_due_diligence import DueDiligencePipeline
    pipeline = DueDiligencePipeline(cim_path="path/to/cim.pdf")
    pipeline.run()
"""

import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Optional

import anthropic

from dd_prompts import PIPELINE, PHASE_NAMES


# ── Document Readers ──────────────────────────────────────────────────

def read_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError:
        sys.exit("pdfplumber required for PDFs: pip install pdfplumber")
    parts = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text()
            if text:
                parts.append(f"[Page {i}]\n{text}")
    return "\n\n".join(parts)


def read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        sys.exit("python-docx required for DOCX: pip install python-docx")
    doc = Document(path)
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def read_document(file_path: str) -> str:
    path = Path(file_path)
    if not path.exists():
        sys.exit(f"File not found: {file_path}")
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8")
    elif suffix == ".pdf":
        return read_pdf(path)
    elif suffix == ".docx":
        return read_docx(path)
    else:
        sys.exit(f"Unsupported file type: {suffix}. Use .pdf, .docx, or .txt")


# ── Variable Store ────────────────────────────────────────────────────

class VariableStore:
    """Holds all intermediate outputs from each pipeline step."""

    def __init__(self):
        self._vars: dict[str, str] = {}

    def set(self, key: str, value: str):
        self._vars[key] = value

    def get(self, key: str, default: str = "[Not yet computed]") -> str:
        return self._vars.get(key, default)

    def substitute(self, template: str) -> str:
        """Replace {{VAR_NAME}} placeholders with stored values."""
        def replacer(match):
            var_name = match.group(1)
            val = self.get(var_name)
            # Truncate very large values to avoid context overflow
            if len(val) > 80_000:
                return val[:80_000] + "\n\n[... truncated for context length ...]"
            return val
        return re.sub(r"\{\{(\w+)\}\}", replacer, template)

    def keys(self):
        return self._vars.keys()

    def save_snapshot(self, output_dir: Path, step_id: str):
        """Save current variable store state to disk."""
        snapshot_path = output_dir / f"step_{step_id.replace('.', '_')}_vars.json"
        # Save a summary (not full text) to avoid huge files
        summary = {}
        for k, v in self._vars.items():
            if len(v) > 5000:
                summary[k] = v[:5000] + f"\n\n[... {len(v)} chars total ...]"
            else:
                summary[k] = v
        snapshot_path.write_text(json.dumps(summary, indent=2, ensure_ascii=False))


# ── Pipeline Engine ───────────────────────────────────────────────────

class DueDiligencePipeline:
    """
    Orchestrates the 8-phase, 30-prompt due diligence pipeline.

    Parameters
    ----------
    cim_path : str
        Path to the CIM document (PDF, DOCX, or TXT).
    deal_terms : dict, optional
        User-provided deal terms (purchase price, leverage, equity split, etc.).
    independent_industry_analysis : str, optional
        JSON string from separate industry research pipeline.
    output_dir : str, optional
        Directory to save all outputs. Defaults to ./dd_output_<timestamp>.
    model : str
        Claude model to use. Defaults to claude-opus-4-6.
    start_phase : int
        Phase to start from (1-8). For resuming interrupted runs.
    start_step : str
        Step ID to start from (e.g., "3.1"). More granular than start_phase.
    """

    def __init__(
        self,
        cim_path: str,
        deal_terms: Optional[dict] = None,
        independent_industry_analysis: Optional[str] = None,
        output_dir: Optional[str] = None,
        model: str = "claude-opus-4-6",
        start_phase: int = 1,
        start_step: Optional[str] = None,
    ):
        self.cim_path = cim_path
        self.model = model
        self.start_phase = start_phase
        self.start_step = start_step
        self.client = anthropic.Anthropic()
        self.store = VariableStore()

        # Set up output directory
        if output_dir:
            self.output_dir = Path(output_dir)
        else:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            self.output_dir = Path(cim_path).parent / f"dd_output_{ts}"
        self.output_dir.mkdir(parents=True, exist_ok=True)

        # Pre-load user-provided inputs
        if deal_terms:
            self.store.set("DEAL_TERMS", json.dumps(deal_terms, indent=2))
        else:
            self.store.set("DEAL_TERMS", json.dumps({
                "purchase_price": "[To be provided by user]",
                "equity_pct": "[To be provided]",
                "debt_pct": "[To be provided]",
                "interest_rate": "[To be provided]",
                "tax_rate": "[To be provided]",
                "exit_multiple": "[To be provided]",
                "hold_period_years": 5
            }, indent=2))

        if independent_industry_analysis:
            self.store.set("INDEPENDENT_INDUSTRY_ANALYSIS", independent_industry_analysis)
        else:
            self.store.set(
                "INDEPENDENT_INDUSTRY_ANALYSIS",
                "[Independent industry analysis not provided. "
                "Using available CIM data with appropriate skepticism toward seller claims. "
                "Flag all market data as requiring independent verification.]"
            )

    def _call_claude(self, system: str, user_content: str, step_name: str) -> str:
        """Make a single Claude API call with streaming."""
        print(f"    Calling Claude ({self.model})...", flush=True)
        start = time.time()

        with self.client.messages.stream(
            model=self.model,
            max_tokens=16384,
            thinking={"type": "adaptive"},
            system=system,
            messages=[{"role": "user", "content": user_content}],
        ) as stream:
            response = stream.get_final_message()

        elapsed = time.time() - start
        input_tokens = response.usage.input_tokens
        output_tokens = response.usage.output_tokens
        print(f"    Done in {elapsed:.1f}s ({input_tokens:,} in / {output_tokens:,} out)")

        # Extract text blocks (skip thinking blocks)
        text_parts = []
        for block in response.content:
            if block.type == "text":
                text_parts.append(block.text)
        return "\n".join(text_parts)

    def _save_step_output(self, step_id: str, step_name: str, output: str):
        """Save a step's raw output to disk."""
        safe_name = re.sub(r"[^\w\-]", "_", f"{step_id}_{step_name}")
        out_path = self.output_dir / f"{safe_name}.md"
        out_path.write_text(output, encoding="utf-8")

    def _should_skip(self, step_id: str, phase: int) -> bool:
        """Check if this step should be skipped based on start_phase/start_step."""
        if self.start_step:
            # Parse step IDs like "3.1" for comparison
            start_parts = [int(x) for x in self.start_step.split(".")]
            step_parts = [int(x) for x in step_id.split(".")]
            return step_parts < start_parts
        return phase < self.start_phase

    def run(self):
        """Execute the full pipeline."""
        print("=" * 70)
        print("  AI DUE DILIGENCE PIPELINE")
        print(f"  CIM: {self.cim_path}")
        print(f"  Output: {self.output_dir}")
        print(f"  Model: {self.model}")
        print(f"  Steps: {len(PIPELINE)}")
        print("=" * 70)

        # Phase 0: Read document
        print("\n[Phase 0] Reading CIM document...")
        cim_text = read_document(self.cim_path)
        self.store.set("CIM_TEXT", cim_text)
        print(f"  Loaded {len(cim_text):,} characters")

        # Save raw CIM text
        (self.output_dir / "00_cim_raw.txt").write_text(cim_text, encoding="utf-8")

        current_phase = 0
        total_steps = len(PIPELINE)
        completed = 0
        errors = []

        for i, prompt in enumerate(PIPELINE):
            step_id = prompt["id"]
            phase = prompt["phase"]
            name = prompt["name"]

            # Phase header
            if phase != current_phase:
                current_phase = phase
                print(f"\n{'=' * 70}")
                print(f"  PHASE {phase}: {PHASE_NAMES[phase].upper()}")
                print(f"{'=' * 70}")

            # Skip check
            if self._should_skip(step_id, phase):
                print(f"\n  [{step_id}] {name} — SKIPPED (resuming from {self.start_step or f'Phase {self.start_phase}'})")
                continue

            print(f"\n  [{step_id}] {name} ({i+1}/{total_steps})")
            print(f"  {'-' * 50}")

            try:
                # Substitute variables into prompts
                system_prompt = self.store.substitute(prompt["system"])
                user_prompt = self.store.substitute(prompt["user"])

                # Call Claude
                output = self._call_claude(system_prompt, user_prompt, name)

                # Store output
                output_var = prompt["output_var"]
                self.store.set(output_var, output)

                # Save to disk
                self._save_step_output(step_id, name, output)

                # Also try to save parsed JSON if output looks like JSON
                self._try_save_json(step_id, name, output)

                completed += 1
                print(f"    -> Stored as {{{{{output_var}}}}}} ({len(output):,} chars)")

            except anthropic.APIError as e:
                error_msg = f"API error at step {step_id}: {e}"
                print(f"    ERROR: {error_msg}")
                errors.append(error_msg)
                # Store error message so downstream steps know
                self.store.set(prompt["output_var"], f"[Error in step {step_id}: {e}]")
                continue

            except Exception as e:
                error_msg = f"Unexpected error at step {step_id}: {e}"
                print(f"    ERROR: {error_msg}")
                errors.append(error_msg)
                self.store.set(prompt["output_var"], f"[Error in step {step_id}: {e}]")
                continue

        # Final summary
        self._print_summary(completed, total_steps, errors)
        self._save_final_report()

        return self.store

    def _try_save_json(self, step_id: str, name: str, output: str):
        """Try to parse and save output as formatted JSON."""
        try:
            # Try to find JSON in the output
            cleaned = output.strip()
            if cleaned.startswith("```"):
                cleaned = cleaned.split("\n", 1)[1]
            if cleaned.endswith("```"):
                cleaned = cleaned.rsplit("```", 1)[0]

            # Try direct parse
            parsed = json.loads(cleaned)
            safe_name = re.sub(r"[^\w\-]", "_", f"{step_id}_{name}")
            json_path = self.output_dir / f"{safe_name}.json"
            json_path.write_text(json.dumps(parsed, indent=2, ensure_ascii=False))
        except (json.JSONDecodeError, ValueError):
            pass  # Not JSON output, that's fine

    def _print_summary(self, completed: int, total: int, errors: list):
        print(f"\n{'=' * 70}")
        print("  PIPELINE COMPLETE")
        print(f"{'=' * 70}")
        print(f"  Steps completed: {completed}/{total}")
        print(f"  Output directory: {self.output_dir}")
        if errors:
            print(f"  Errors ({len(errors)}):")
            for e in errors:
                print(f"    - {e}")
        print()

    def _save_final_report(self):
        """Assemble all outputs into a single markdown report."""
        report_parts = [
            "# AI Due Diligence Report",
            f"**Generated**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"**CIM**: {self.cim_path}",
            f"**Model**: {self.model}",
            "",
            "---",
            "",
        ]

        current_phase = 0
        for prompt in PIPELINE:
            phase = prompt["phase"]
            if phase != current_phase:
                current_phase = phase
                report_parts.append(f"\n## Phase {phase}: {PHASE_NAMES[phase]}\n")

            output_var = prompt["output_var"]
            output = self.store.get(output_var, "")
            if output and not output.startswith("["):
                report_parts.append(f"### {prompt['id']} – {prompt['name']}\n")
                report_parts.append(output)
                report_parts.append("\n---\n")

        report = "\n".join(report_parts)
        report_path = self.output_dir / "FULL_IC_REPORT.md"
        report_path.write_text(report, encoding="utf-8")
        print(f"  Full report saved: {report_path}")

    def run_single_step(self, step_id: str) -> str:
        """Run a single step by its ID (e.g., '3.1'). Useful for testing."""
        prompt = next((p for p in PIPELINE if p["id"] == step_id), None)
        if not prompt:
            raise ValueError(f"Step {step_id} not found. Valid: {[p['id'] for p in PIPELINE]}")

        # Load CIM if not already loaded
        if not self.store.get("CIM_TEXT", "").strip() or self.store.get("CIM_TEXT") == "[Not yet computed]":
            cim_text = read_document(self.cim_path)
            self.store.set("CIM_TEXT", cim_text)

        system_prompt = self.store.substitute(prompt["system"])
        user_prompt = self.store.substitute(prompt["user"])

        output = self._call_claude(system_prompt, user_prompt, prompt["name"])
        self.store.set(prompt["output_var"], output)
        self._save_step_output(step_id, prompt["name"], output)

        return output


# ── Convenience function ──────────────────────────────────────────────

def run_pipeline(
    cim_path: str,
    deal_terms: Optional[dict] = None,
    independent_industry_analysis: Optional[str] = None,
    output_dir: Optional[str] = None,
    model: str = "claude-opus-4-6",
    start_phase: int = 1,
    start_step: Optional[str] = None,
) -> VariableStore:
    """Convenience function to run the full pipeline."""
    pipeline = DueDiligencePipeline(
        cim_path=cim_path,
        deal_terms=deal_terms,
        independent_industry_analysis=independent_industry_analysis,
        output_dir=output_dir,
        model=model,
        start_phase=start_phase,
        start_step=start_step,
    )
    return pipeline.run()
